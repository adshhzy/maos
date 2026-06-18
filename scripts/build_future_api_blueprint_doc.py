from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path.home() / "Desktop" / "持久化多Agent编排内核_未来完整API蓝图.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5F6B7A"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "B8C2CF"
LIGHT_GRAY = "F7F9FC"
WARNING_FILL = "FFF7E6"


def main() -> None:
    doc = Document()
    configure_document(doc)

    add_title(
        doc,
        "持久化多 Agent 编排内核未来完整 API 蓝图",
        "面向 Temporal + Google A2A、真实 Agent Service、人机协作和大规模长等待任务的接口规划",
    )
    add_metadata(doc)
    add_callout(
        doc,
        "定位",
        "持久化多 Agent 编排内核只负责用户任务的分工、编排、调度、状态持久化和唤醒；"
        "每个图节点是一轮 Agent 调用，实际工作由外部 Agent Service 或人类参与者完成。"
        "Temporal 承担耐久执行、长等待、重试、信号唤醒和可恢复历史；A2A 承担 Agent 间消息、任务、状态和 Artifact 的协议形态。",
    )

    add_heading(doc, "1. API 总览", 1)
    add_paragraph(
        doc,
        "当前实现已经具备任务创建、批量运行、Temporal 持久 workflow、模拟 Agent 回调和 Web 可视化。"
        "未来完整体应把接口分为控制面、Agent 数据面、事件面、Artifact 面、人机协作面和运维面。"
        "其中面向用户和 Web 的接口建议统一升级为 /api/v1；面向 Agent Service 的接口应优先采用 A2A 标准方法，"
        "只在容量、租户、治理等 A2A 未覆盖的地方增加 MAOS 扩展接口。",
    )
    add_table(
        doc,
        ["接口层", "主要调用方", "职责边界", "代表接口"],
        [
            [
                "任务控制面",
                "Web UI、CLI、上层业务系统",
                "创建、查询、暂停、恢复、取消、重试和终止 DAG 任务。",
                "POST /api/v1/tasks；GET /api/v1/tasks/{taskId}",
            ],
            [
                "DAG 模板面",
                "平台管理员、任务设计器",
                "校验、版本化、发布和实例化任务图。",
                "POST /api/v1/graph-templates:validate",
            ],
            [
                "人机协作面",
                "人类参与者、审批系统、客服/运营后台",
                "处理需要人工输入、审批、确认、补充材料或冲突裁决的节点。",
                "POST /api/v1/human-tasks/{humanTaskId}:submit",
            ],
            [
                "Agent 注册与治理面",
                "Agent Service、平台运维",
                "登记 Agent 能力、健康、容量、租户范围和上线下线状态。",
                "POST /api/v1/agents；PATCH /api/v1/agents/{agentId}/capacity",
            ],
            [
                "A2A Agent 数据面",
                "编排内核 -> Agent Service",
                "把节点输入、依赖结果和上下文作为 A2A Message/Task/Artifact 交给 Agent。",
                "SendMessage；SubscribeToTask；GetTask",
            ],
            [
                "Agent 事件回调面",
                "Agent Service -> 编排内核",
                "报告 accepted、working、artifact、input_required、completed、failed 等事件，并唤醒 Temporal workflow。",
                "POST /api/v1/agent-events",
            ],
            [
                "Artifact 面",
                "Agent、人类、Web UI、外部存储",
                "上传、下载、引用、权限、血缘和生命周期管理。Temporal 只保存 ArtifactRef。",
                "POST /api/v1/artifacts；GET /api/v1/artifacts/{artifactId}",
            ],
            [
                "观测与运维面",
                "运维、SRE、看板、审计系统",
                "健康、指标、事件流、死信、补偿、对账、Temporal 集群状态。",
                "GET /api/v1/metrics；POST /api/v1/reconcile",
            ],
        ],
        widths=[1.35, 1.45, 2.15, 1.55],
    )

    add_heading(doc, "2. 通用约定", 1)
    add_table(
        doc,
        ["约定", "建议"],
        [
            ["Base URL", "用户侧和管理侧统一使用 /api/v1；旧的 /api/tasks、/api/agent-callbacks 可作为兼容层保留。"],
            ["认证", "Authorization: Bearer <token>；服务间可用 mTLS 或短期 JWT。"],
            ["租户", "X-MAOS-Tenant 或 token claim 中的 tenant_id；所有查询、事件、Artifact 均按租户隔离。"],
            ["幂等", "写接口要求 Idempotency-Key；Agent 事件要求 event_id + agent_task_id + sequence 去重。"],
            ["追踪", "X-MAOS-Trace-Id / correlation_id 贯穿 Web、Sandbox、Temporal、Agent Service 和 Artifact Store。"],
            ["回调签名", "Agent -> Sandbox 回调使用 X-MAOS-Signature，签名体包含时间戳、event_id、agent_task_id、payload hash。"],
            ["长等待", "节点派发后 workflow 进入 Temporal 等待；Agent 完成后通过 callback/push event 转为 workflow signal。"],
            ["大结果", "不把大文本、文件或二进制塞进 Temporal history；只传 ArtifactRef、摘要、校验和和元数据。"],
        ],
        widths=[1.55, 4.95],
    )
    add_table(
        doc,
        ["ID", "含义"],
        [
            ["tenant_id", "租户或组织隔离边界。"],
            ["task_id / workflow_id", "用户任务实例 ID；建议直接等于 Temporal workflow_id。"],
            ["graph_id / graph_version", "DAG 模板或一次性 DAG 的业务标识与版本。"],
            ["node_id", "DAG 中的节点 ID；一个节点代表一次 Agent 调用或人工任务。"],
            ["agent_id", "Agent Service 在注册中心中的稳定 ID。"],
            ["agent_task_id", "Agent Service 内部的 A2A Task ID；由 SendMessage 返回或事件中携带。"],
            ["artifact_id", "Artifact Store 中的结果、输入材料或中间产物 ID。"],
            ["human_task_id", "人机协作子任务 ID。"],
            ["event_id", "Agent 事件、用户事件或系统事件的幂等 ID。"],
        ],
        widths=[1.75, 4.75],
    )

    add_heading(doc, "3. 任务控制面 API", 1)
    add_paragraph(
        doc,
        "这些接口是 Web UI、CLI 和上层业务系统最常用的入口。它们不直接持有任务执行线程，"
        "而是通过 Temporal workflow handle 查询、signal、cancel 或 terminate 实例。",
    )
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["POST /api/v1/tasks", "Client -> Sandbox", "创建单个 DAG 任务。", "graph、input、priority、labels、deadline、idempotency_key。"],
            ["POST /api/v1/tasks:batchCreate", "Client -> Sandbox", "批量创建任务图并并发启动。", "graphs[]；返回 task_ids[] 和逐项错误。"],
            ["GET /api/v1/tasks", "Client -> Sandbox", "分页查询任务列表。", "status、agent_id、graph_id、created_after、cursor、limit。"],
            ["GET /api/v1/tasks/{taskId}", "Client -> Sandbox", "查询任务详情。", "汇总状态、Temporal 状态、节点状态、最近事件、Artifact 摘要。"],
            ["PATCH /api/v1/tasks/{taskId}", "Client -> Sandbox", "更新任务元数据。", "name、labels、priority、deadline、visibility。"],
            ["POST /api/v1/tasks/{taskId}:pause", "Client -> Sandbox", "请求业务级暂停。", "通过 workflow signal 设置暂停门禁；不杀 workflow。"],
            ["POST /api/v1/tasks/{taskId}:resume", "Client -> Sandbox", "恢复业务级暂停任务。", "解除暂停门禁并继续调度 ready 节点。"],
            ["POST /api/v1/tasks/{taskId}:cancel", "Client -> Sandbox", "请求可恢复语义的取消。", "通知未完成 Agent Task CancelTask；workflow 走取消路径。"],
            ["POST /api/v1/tasks/{taskId}:terminate", "Admin -> Sandbox", "强制终止不可恢复任务。", "映射 Temporal terminate；需要管理员权限和 reason。"],
            ["POST /api/v1/tasks/{taskId}:retry", "Client -> Sandbox", "按策略重跑失败任务。", "retry_scope: failed_nodes | whole_workflow；preserve_artifacts。"],
            ["GET /api/v1/tasks/{taskId}/graph", "Client -> Sandbox", "获取可视化 DAG。", "nodes、edges、levels、dependency_status。"],
            ["GET /api/v1/tasks/{taskId}/timeline", "Client -> Sandbox", "获取任务时间线。", "派发、挂起、唤醒、Artifact、人类输入、失败事件。"],
            ["GET /api/v1/tasks/{taskId}/events", "Client -> Sandbox", "按 cursor 拉取任务事件。", "cursor、limit、event_type、since。"],
            ["GET /api/v1/tasks/{taskId}/artifacts", "Client -> Sandbox", "列出任务所有产物。", "按 node_id、artifact_type、visibility 过滤。"],
        ],
        widths=[1.85, 1.25, 1.8, 1.6],
    )

    add_heading(doc, "4. 节点控制面 API", 1)
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["GET /api/v1/tasks/{taskId}/nodes", "Client -> Sandbox", "列出节点状态。", "node_id、state、attempt、agent_id、agent_task_id、dependencies。"],
            ["GET /api/v1/tasks/{taskId}/nodes/{nodeId}", "Client -> Sandbox", "查询单个节点详情。", "输入、依赖 ArtifactRef、Agent 事件、错误、耗时。"],
            ["POST /api/v1/tasks/{taskId}/nodes/{nodeId}:retry", "Client -> Sandbox", "重试失败或被取消节点。", "attempt_policy、reuse_dependency_outputs。"],
            ["POST /api/v1/tasks/{taskId}/nodes/{nodeId}:skip", "Admin -> Sandbox", "跳过节点并可选提供替代结果。", "override_artifacts、reason、audit_actor。"],
            ["POST /api/v1/tasks/{taskId}/nodes/{nodeId}:cancel", "Client -> Sandbox", "取消正在执行的 Agent 节点。", "向 Agent Service 调 CancelTask，并 signal workflow。"],
            ["POST /api/v1/tasks/{taskId}/nodes/{nodeId}:overrideResult", "Admin -> Sandbox", "人工覆盖节点结果。", "result_summary、artifact_refs、downstream_policy。"],
            ["GET /api/v1/tasks/{taskId}/nodes/{nodeId}/events", "Client -> Sandbox", "节点级事件流查询。", "只返回该节点的 Agent、人类和系统事件。"],
            ["GET /api/v1/tasks/{taskId}/nodes/{nodeId}/artifacts", "Client -> Sandbox", "节点产物列表。", "output_artifacts、debug_artifacts、dependency_artifacts。"],
        ],
        widths=[2.05, 1.15, 1.75, 1.55],
    )

    add_heading(doc, "5. DAG 模板与任务图 API", 1)
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["POST /api/v1/graph-templates", "Designer -> Sandbox", "创建 DAG 模板。", "template_id、version、schema、nodes、edges、policies。"],
            ["GET /api/v1/graph-templates", "Designer -> Sandbox", "分页查询模板。", "owner、status、tag、capability、cursor。"],
            ["GET /api/v1/graph-templates/{templateId}", "Designer -> Sandbox", "查看模板最新版本。", "包含版本列表和发布状态。"],
            ["POST /api/v1/graph-templates/{templateId}/versions", "Designer -> Sandbox", "提交新版本。", "graph_spec、migration_notes、compatibility。"],
            ["POST /api/v1/graph-templates:validate", "Designer -> Sandbox", "校验任意 DAG JSON。", "检查 DAG 无环、依赖存在、Agent 能力匹配、输入输出 schema。"],
            ["POST /api/v1/graph-templates/{templateId}:instantiate", "Client -> Sandbox", "由模板创建任务实例。", "template_version、input、runtime_overrides。"],
            ["POST /api/v1/graphs:dryRun", "Designer -> Sandbox", "静态演练调度计划。", "返回拓扑层级、并行度、预估 Agent 能力缺口。"],
            ["POST /api/v1/graphs:estimate", "Designer -> Sandbox", "估算资源与时长。", "基于历史耗时、Agent 容量、人工等待策略。"],
        ],
        widths=[2.1, 1.15, 1.75, 1.5],
    )
    add_paragraph(doc, "建议的 DAG 节点定义至少包含：")
    add_code(
        doc,
        "{\n"
        '  "id": "review_contract",\n'
        '  "type": "agent_task",\n'
        '  "agent": {"required_skill": "contract.review", "preferred_agent_id": "legal-agent"},\n'
        '  "depends_on": ["extract_terms", "fetch_policy"],\n'
        '  "input": {"prompt": "审查合同风险", "schema": "ContractReviewInput"},\n'
        '  "output": {"schema": "ContractReviewResult", "artifact_types": ["risk_report"]},\n'
        '  "policies": {"timeout": "24h", "retry": {"max_attempts": 3}, "human_in_loop": "on_input_required"}\n'
        "}",
    )

    add_heading(doc, "6. 人机协作 API", 1)
    add_paragraph(
        doc,
        "在多 Agent + 多人协作场景中，Agent 可能需要用户补充信息、人类审批、专家裁决或外部系统确认。"
        "这些等待也应通过 Temporal signal 或 update 唤醒 workflow，而不是由管理层长期轮询。",
    )
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["GET /api/v1/human-tasks", "Human UI -> Sandbox", "列出待办。", "assignee、role、task_id、node_id、status、due_before。"],
            ["GET /api/v1/human-tasks/{humanTaskId}", "Human UI -> Sandbox", "查看待办详情。", "问题、上下文、依赖 Artifact、允许动作、审计历史。"],
            ["POST /api/v1/human-tasks/{humanTaskId}:claim", "Human UI -> Sandbox", "领取待办。", "actor_id、claim_ttl、conflict_policy。"],
            ["POST /api/v1/human-tasks/{humanTaskId}:submit", "Human UI -> Sandbox", "提交人工输入。", "answer、artifact_refs、decision、comment；触发 workflow signal。"],
            ["POST /api/v1/human-tasks/{humanTaskId}:reject", "Human UI -> Sandbox", "拒绝或退回。", "reason、return_to_agent、required_changes。"],
            ["POST /api/v1/human-tasks/{humanTaskId}:delegate", "Human UI -> Sandbox", "转派。", "target_user、target_role、reason。"],
            ["POST /api/v1/human-tasks/{humanTaskId}:escalate", "System -> Sandbox", "超时升级。", "escalation_policy、new_due_at、notification_channels。"],
            ["POST /api/v1/tasks/{taskId}/human-events", "External -> Sandbox", "外部人工系统事件入口。", "event_id、human_task_id、action、payload。"],
        ],
        widths=[2.1, 1.15, 1.75, 1.5],
    )

    add_heading(doc, "7. Agent 注册与治理 API", 1)
    add_paragraph(
        doc,
        "A2A 负责 Agent 通信协议，但编排内核仍需要自己的 Agent 注册与治理层，用于路由、租户隔离、容量、限流、健康检查、灰度和审计。"
    )
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["POST /api/v1/agents", "Agent Service -> Sandbox", "注册 Agent Service。", "agent_id、base_url、tenant_scope、skills、auth、agent_card_url。"],
            ["GET /api/v1/agents", "Admin -> Sandbox", "列出 Agent。", "status、skill、tenant、version、cursor。"],
            ["GET /api/v1/agents/{agentId}", "Admin -> Sandbox", "查看 Agent 详情。", "AgentCard、能力、容量、最近心跳、错误率。"],
            ["PATCH /api/v1/agents/{agentId}", "Admin -> Sandbox", "更新注册信息。", "base_url、labels、routing_weight、auth_profile。"],
            ["DELETE /api/v1/agents/{agentId}", "Admin -> Sandbox", "删除或注销 Agent。", "仅允许无运行中任务或强制 drain 后执行。"],
            ["POST /api/v1/agents/{agentId}:heartbeat", "Agent Service -> Sandbox", "主动心跳。", "load、running_tasks、queue_depth、available_slots、build_version。"],
            ["POST /api/v1/agents/{agentId}:drain", "Admin -> Sandbox", "停止接新任务并等待清空。", "drain_mode、deadline、replacement_agent_id。"],
            ["POST /api/v1/agents/{agentId}:enable", "Admin -> Sandbox", "启用路由。", "允许新节点调度到该 Agent。"],
            ["POST /api/v1/agents/{agentId}:disable", "Admin -> Sandbox", "禁用路由。", "不再分配新任务，已运行任务按策略处理。"],
            ["GET /api/v1/agents/{agentId}/capacity", "Scheduler -> Sandbox", "查询容量。", "max_concurrency、available_slots、rate_limits、tenant_quota。"],
            ["PATCH /api/v1/agents/{agentId}/capacity", "Admin -> Sandbox", "调整容量与限流。", "per_tenant、per_skill、burst、cooldown。"],
            ["GET /api/v1/agents/{agentId}/card", "Sandbox -> Agent", "获取缓存 AgentCard。", "可代理 A2A /.well-known/agent-card.json。"],
            ["POST /api/v1/agents/{agentId}/card:refresh", "Sandbox -> Agent", "刷新 AgentCard。", "重新拉取 public/extended card 并校验签名。"],
        ],
        widths=[2.1, 1.15, 1.75, 1.5],
    )

    add_heading(doc, "8. 编排内核 -> Agent Service：A2A 标准 API", 1)
    add_paragraph(
        doc,
        "真实 Agent Service 应尽量实现 A2A Server。编排内核作为 A2A Client，把每个 DAG 节点转换为一次 SendMessage 或 SendStreamingMessage。"
        "当前 simulator 的 POST /api/simulator/jobs 只是临时替代物，未来应被这一组 A2A 接口取代。",
    )
    add_table(
        doc,
        ["A2A 方法 / REST", "方向", "用途", "编排内核使用方式"],
        [
            ["GET /.well-known/agent-card.json", "Sandbox -> Agent", "发现 Agent 能力、协议、认证、技能和扩展。", "注册或刷新 Agent 时读取，用于路由和能力匹配。"],
            ["GET /extendedAgentCard / GetExtendedAgentCard", "Sandbox -> Agent", "认证后获取更完整能力信息。", "读取私有技能、限额、模型、工具、输入输出 schema。"],
            ["POST /message:send / SendMessage", "Sandbox -> Agent", "发起或继续一个 Agent Task。", "派发节点；configuration.return_immediately 建议为 true。"],
            ["POST /message:stream / SendStreamingMessage", "Sandbox -> Agent", "发起任务并接收实时流。", "用于需要实时 UI 展示或短任务同步反馈的场景。"],
            ["GET /tasks/{id} / GetTask", "Sandbox -> Agent", "查询 Agent Task 当前状态。", "仅用于恢复、对账或回调丢失补偿，不作为主轮询路径。"],
            ["GET /tasks / ListTasks", "Sandbox -> Agent", "分页查询 Agent 侧任务。", "运维对账、Agent 重启恢复、租户审计。"],
            ["POST /tasks/{id}:cancel / CancelTask", "Sandbox -> Agent", "请求取消 Agent Task。", "用户取消节点、workflow 取消或超时策略触发。"],
            ["POST /tasks/{id}:subscribe / SubscribeToTask", "Sandbox -> Agent", "订阅已有任务事件流。", "用于恢复事件流或临时调试；长期生产更推荐 push notification。"],
            ["POST /tasks/{id}/pushNotificationConfigs / CreateTaskPushNotificationConfig", "Sandbox -> Agent", "为 Agent Task 配置回调地址。", "把 /api/v1/agent-events 作为 push endpoint。"],
            ["GET /tasks/{id}/pushNotificationConfigs/{configId} / GetTaskPushNotificationConfig", "Sandbox -> Agent", "查询回调配置。", "对账和调试。"],
            ["GET /tasks/{id}/pushNotificationConfigs / ListTaskPushNotificationConfigs", "Sandbox -> Agent", "列出回调配置。", "确认某个任务是否已配置回调。"],
            ["DELETE /tasks/{id}/pushNotificationConfigs/{configId} / DeleteTaskPushNotificationConfig", "Sandbox -> Agent", "删除回调配置。", "任务结束、取消或租户权限变更时清理。"],
        ],
        widths=[2.05, 1.2, 1.65, 1.6],
    )
    add_paragraph(doc, "SendMessage 的节点派发请求建议携带如下元数据：")
    add_code(
        doc,
        "{\n"
        '  "message": {\n'
        '    "role": "user",\n'
        '    "parts": [{"kind": "text", "text": "请执行节点 review_contract"}],\n'
        '    "contextId": "task-20260612-001",\n'
        '    "referenceTaskIds": ["a2a-task-extract_terms", "a2a-task-fetch_policy"],\n'
        '    "metadata": {\n'
        '      "workflow_id": "task-20260612-001",\n'
        '      "node_id": "review_contract",\n'
        '      "attempt": 1,\n'
        '      "dependency_artifacts": [{"artifact_id": "art-terms", "producer_node_id": "extract_terms"}],\n'
        '      "deadline": "2026-06-13T08:00:00Z",\n'
        '      "callback_event_url": "https://sandbox.example.com/api/v1/agent-events"\n'
        "    }\n"
        "  },\n"
        '  "configuration": {\n'
        '    "return_immediately": true,\n'
        '    "pushNotificationConfig": {"url": "https://sandbox.example.com/api/v1/agent-events"}\n'
        "  }\n"
        "}",
    )

    add_heading(doc, "9. Agent Service -> 编排内核：事件与回调 API", 1)
    add_paragraph(
        doc,
        "Agent 完成、失败、需要人工输入或产生 Artifact 时，应主动推送事件给编排内核。"
        "编排内核收到事件后做幂等落库、ArtifactRef 校验，然后 signal 对应 Temporal workflow。"
    )
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["POST /api/v1/agent-events", "Agent -> Sandbox", "统一事件入口。", "event_id、workflow_id、node_id、agent_task_id、event_type、sequence、payload。"],
            ["POST /api/v1/agent-events/status", "Agent -> Sandbox", "状态事件快捷入口。", "accepted、working、suspended、input_required、completed、failed、canceled。"],
            ["POST /api/v1/agent-events/artifact", "Agent -> Sandbox", "Artifact 增量事件。", "TaskArtifactUpdateEvent、artifact_ref、append、last_chunk。"],
            ["POST /api/v1/agent-events/completed", "Agent -> Sandbox", "终态成功事件。", "final_artifacts、result_summary、usage、metrics。"],
            ["POST /api/v1/agent-events/failed", "Agent -> Sandbox", "终态失败事件。", "error_code、retryable、error_message、diagnostics_artifact。"],
            ["POST /api/v1/agent-events/input-required", "Agent -> Sandbox", "请求人工输入。", "question、choices、required_artifacts、human_task_policy。"],
            ["POST /api/v1/agent-events/heartbeat", "Agent -> Sandbox", "节点级心跳。", "progress、stage、estimated_remaining、resource_usage。"],
            ["POST /api/v1/a2a/tasks/{agentTaskId}/events", "Agent -> Sandbox", "A2A 事件兼容入口。", "直接接收 TaskStatusUpdateEvent / TaskArtifactUpdateEvent 包装。"],
        ],
        widths=[2.05, 1.15, 1.75, 1.55],
    )
    add_code(
        doc,
        "{\n"
        '  "event_id": "evt-01HX...",\n'
        '  "workflow_id": "task-20260612-001",\n'
        '  "node_id": "review_contract",\n'
        '  "agent_id": "legal-agent",\n'
        '  "agent_task_id": "a2a-task-review_contract-001",\n'
        '  "attempt": 1,\n'
        '  "sequence": 7,\n'
        '  "event_type": "completed",\n'
        '  "occurred_at": "2026-06-12T08:30:00Z",\n'
        '  "a2a": {"statusUpdate": {"status": {"state": "completed"}}},\n'
        '  "artifact_refs": [{"artifact_id": "art-risk-report", "kind": "risk_report"}],\n'
        '  "result_summary": {"risk_level": "medium"}\n'
        "}",
    )

    add_heading(doc, "10. Artifact API", 1)
    add_paragraph(
        doc,
        "Artifact 是 Agent 之间传递大结果和上下文的主要媒介。Temporal history 中只保留 ArtifactRef、摘要和校验信息；"
        "真正内容进入对象存储、数据库或外部文档系统。",
    )
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["POST /api/v1/artifacts", "Agent/Human -> Sandbox", "创建 Artifact 元数据或申请上传。", "kind、media_type、size、sha256、producer、visibility。"],
            ["POST /api/v1/artifacts/{artifactId}:completeMultipart", "Agent -> Sandbox", "完成分片上传。", "parts、checksum、final_size。"],
            ["POST /api/v1/artifacts/{artifactId}:signUpload", "Agent/Human -> Sandbox", "获取上传签名 URL。", "content_type、expires_in、max_size。"],
            ["POST /api/v1/artifacts/{artifactId}:signDownload", "Client/Agent -> Sandbox", "获取下载签名 URL。", "scope、expires_in、watermark_policy。"],
            ["GET /api/v1/artifacts/{artifactId}", "Client/Agent -> Sandbox", "读取 Artifact 元数据。", "不直接返回大文件内容。"],
            ["GET /api/v1/tasks/{taskId}/artifacts", "Client -> Sandbox", "按任务列出 Artifact。", "node_id、kind、producer_type、cursor。"],
            ["GET /api/v1/tasks/{taskId}/nodes/{nodeId}/artifacts", "Client -> Sandbox", "按节点列出 Artifact。", "输入依赖、输出、调试材料分组。"],
            ["POST /api/v1/artifacts/{artifactId}:link", "Sandbox -> Sandbox", "把已有 Artifact 链接给节点。", "target_task_id、target_node_id、relation。"],
            ["DELETE /api/v1/artifacts/{artifactId}", "Admin -> Sandbox", "删除或标记过期。", "retention_policy、legal_hold、audit_reason。"],
        ],
        widths=[2.1, 1.15, 1.75, 1.5],
    )
    add_code(
        doc,
        "{\n"
        '  "artifact_id": "art-risk-report",\n'
        '  "kind": "risk_report",\n'
        '  "media_type": "application/json",\n'
        '  "uri": "s3://maos-artifacts/tenant-a/task-001/art-risk-report.json",\n'
        '  "sha256": "abc...",\n'
        '  "producer": {"task_id": "task-001", "node_id": "review_contract", "agent_task_id": "a2a-task-001"},\n'
        '  "summary": {"risk_level": "medium"},\n'
        '  "created_at": "2026-06-12T08:30:00Z"\n'
        "}",
    )

    add_heading(doc, "11. 事件、订阅与看板 API", 1)
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["GET /api/v1/events", "Dashboard -> Sandbox", "全局事件分页拉取。", "cursor、tenant、event_type、since、limit。"],
            ["GET /api/v1/tasks/{taskId}/events", "Dashboard -> Sandbox", "任务事件分页拉取。", "适合 30 秒刷新和手动刷新按钮。"],
            ["GET /api/v1/tasks/{taskId}/stream", "Dashboard -> Sandbox", "单任务 SSE 事件流。", "可选能力；不作为默认高频看板依赖。"],
            ["GET /api/v1/dashboard/tasks", "Dashboard -> Sandbox", "所有任务看板快照。", "状态分组、正在执行节点、挂起原因、Agent 分布、错误摘要。"],
            ["GET /api/v1/dashboard/agents", "Dashboard -> Sandbox", "Agent 看板快照。", "容量、负载、健康、running_tasks、queue_depth。"],
            ["POST /api/v1/webhooks/subscriptions", "External -> Sandbox", "创建事件推送订阅。", "url、event_types、filters、signature_secret。"],
            ["GET /api/v1/webhooks/subscriptions", "External -> Sandbox", "列出订阅。", "租户可见范围。"],
            ["DELETE /api/v1/webhooks/subscriptions/{id}", "External -> Sandbox", "删除订阅。", "立即停止后续推送。"],
        ],
        widths=[2.05, 1.15, 1.75, 1.55],
    )

    add_heading(doc, "12. 运维、治理与恢复 API", 1)
    add_table(
        doc,
        ["接口", "方向", "用途", "关键字段 / 语义"],
        [
            ["GET /api/v1/health", "Ops -> Sandbox", "进程级健康检查。", "服务存活、版本、基本依赖。"],
            ["GET /api/v1/ready", "Ops -> Sandbox", "流量就绪检查。", "Temporal 连接、worker、数据库、Artifact Store、Agent Registry。"],
            ["GET /api/v1/metrics", "Prometheus -> Sandbox", "指标。", "任务数、挂起数、失败率、Agent 延迟、callback 延迟。"],
            ["GET /api/v1/runtime", "Ops -> Sandbox", "运行时信息。", "Temporal namespace、task_queue、build、配置摘要。"],
            ["GET /api/v1/temporal", "Ops -> Sandbox", "Temporal 依赖状态。", "server target、UI、namespace、worker poller。"],
            ["POST /api/v1/reconcile", "Ops -> Sandbox", "全局对账。", "检查 Temporal、Agent Task、Artifact、事件表一致性。"],
            ["POST /api/v1/tasks/{taskId}:reconcile", "Ops -> Sandbox", "单任务对账。", "补偿丢失回调、重新订阅 Agent Task、重建看板快照。"],
            ["GET /api/v1/dead-letter-events", "Ops -> Sandbox", "查看死信事件。", "签名失败、未知 workflow、乱序不可恢复、schema 错误。"],
            ["POST /api/v1/dead-letter-events/{eventId}:replay", "Ops -> Sandbox", "重放死信事件。", "修复配置或人工确认后重放。"],
            ["GET /api/v1/audit-logs", "Security -> Sandbox", "审计日志。", "谁在何时创建、取消、覆盖、重试、下载 Artifact。"],
            ["GET /api/v1/quotas", "Ops -> Sandbox", "查询租户和 Agent 配额。", "并发、QPS、Artifact 存储、人工任务限制。"],
            ["PATCH /api/v1/quotas/{scope}", "Ops -> Sandbox", "调整配额。", "tenant、agent、skill、priority 维度。"],
        ],
        widths=[2.05, 1.15, 1.75, 1.55],
    )

    add_heading(doc, "13. 核心对象 Schema 摘要", 1)
    add_heading(doc, "13.1 CreateTaskRequest", 2)
    add_code(
        doc,
        "{\n"
        '  "graph": {"id": "contract-review", "nodes": [], "edges": []},\n'
        '  "input": {"customer_id": "c-001", "documents": [{"artifact_id": "art-contract"}]},\n'
        '  "options": {"priority": "normal", "deadline": "2026-06-13T08:00:00Z"},\n'
        '  "labels": {"project": "legal-ai"},\n'
        '  "idempotency_key": "client-generated-key"\n'
        "}",
    )
    add_heading(doc, "13.2 TaskDetail / NodeState", 2)
    add_code(
        doc,
        "{\n"
        '  "task_id": "task-001",\n'
        '  "workflow_id": "task-001",\n'
        '  "status": "running",\n'
        '  "workflow_state": "waiting_agent",\n'
        '  "suspended": true,\n'
        '  "current_nodes": ["review_contract"],\n'
        '  "nodes": [{\n'
        '    "node_id": "review_contract",\n'
        '    "state": "suspended",\n'
        '    "agent_id": "legal-agent",\n'
        '    "agent_task_id": "a2a-task-review_contract-001",\n'
        '    "attempt": 1,\n'
        '    "dependencies": ["extract_terms"],\n'
        '    "input_artifacts": [{"artifact_id": "art-terms"}],\n'
        '    "output_artifacts": []\n'
        "  }]\n"
        "}",
    )
    add_heading(doc, "13.3 AgentRegistration / HumanTask", 2)
    add_code(
        doc,
        "{\n"
        '  "agent": {\n'
        '    "agent_id": "legal-agent",\n'
        '    "base_url": "https://legal-agent.example.com",\n'
        '    "agent_card_url": "https://legal-agent.example.com/.well-known/agent-card.json",\n'
        '    "skills": ["contract.review", "risk.analysis"],\n'
        '    "capacity": {"max_concurrency": 50, "tenant_limits": {"tenant-a": 10}}\n'
        "  },\n"
        '  "human_task": {\n'
        '    "human_task_id": "ht-001",\n'
        '    "task_id": "task-001",\n'
        '    "node_id": "review_contract",\n'
        '    "status": "open",\n'
        '    "question": "请确认高风险条款是否接受",\n'
        '    "allowed_actions": ["approve", "reject", "request_changes"]\n'
        "  }\n"
        "}",
    )

    add_heading(doc, "14. 典型交互流程", 1)
    add_table(
        doc,
        ["步骤", "参与方", "接口 / Temporal 动作", "状态变化"],
        [
            ["1", "Client -> Sandbox", "POST /api/v1/tasks", "创建 task_id，并 start_workflow(JsonDagWorkflow.run)。"],
            ["2", "Workflow", "计算 ready nodes", "并行派发无依赖或依赖已完成节点。"],
            ["3", "Sandbox -> Agent", "SendMessage 或 SendStreamingMessage", "节点进入 dispatched / waiting_agent。"],
            ["4", "Workflow", "workflow.wait_condition 等待 agent_node_completed signal", "workflow 可被 Temporal 持久化挂起，不占用 worker 线程。"],
            ["5", "Agent -> Sandbox", "POST /api/v1/agent-events", "Sandbox 幂等记录事件，校验 ArtifactRef。"],
            ["6", "Sandbox -> Temporal", "handle.signal(agent_node_completed, event)", "workflow 被唤醒，写入节点结果。"],
            ["7", "Workflow", "解锁下游节点", "继续调度并行节点或进入完成状态。"],
            ["8", "Dashboard -> Sandbox", "GET /api/v1/dashboard/tasks", "30 秒刷新或手动刷新，不需要常驻监听所有 workflow。"],
        ],
        widths=[0.55, 1.35, 2.3, 2.3],
    )

    add_heading(doc, "15. 与当前实现的演进关系", 1)
    add_table(
        doc,
        ["阶段", "目标", "主要改动"],
        [
            ["Phase 1", "接口版本化", "保留旧接口，新增 /api/v1/tasks、/api/v1/agent-events、/api/v1/dashboard/tasks。"],
            ["Phase 2", "真实 A2A Agent 接入", "用 AgentCard 发现替换手写 simulator 协议；SendMessage 替代 POST /api/simulator/jobs。"],
            ["Phase 3", "ArtifactRef 外置存储", "结果从 Temporal payload 迁出，只保存 ArtifactRef、摘要、hash 和 schema。"],
            ["Phase 4", "人机协作闭环", "引入 HumanTask API，把 input_required 事件映射为人工待办和 workflow signal。"],
            ["Phase 5", "治理与容量", "Agent Registry、capacity、quota、drain、enable/disable、租户隔离和路由权重。"],
            ["Phase 6", "生产级恢复", "reconcile、dead-letter、event replay、audit、metrics、webhook subscription。"],
        ],
        widths=[0.85, 1.65, 4.0],
    )

    add_callout(
        doc,
        "关键设计取舍",
        "任务管理层不再维护常驻内存任务实例，而是把每个用户任务作为独立 Temporal workflow。"
        "Web 看板通过分页查询、Temporal visibility/query、事件表和快照接口按需读取状态；Agent 完成时通过 A2A push/callback 触发 signal。"
        "这样大规模并发、长时间等待和人机回路都复用 Temporal 的耐久等待机制，同时把真实执行能力留给外部 Agent Service。",
        fill=WARNING_FILL,
    )

    add_footer(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    set_style_font(styles["Normal"], "Calibri", 11, INK, east_asia="Microsoft YaHei")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        set_style_font(styles[name], "Calibri", size, color, bold=True, east_asia="Microsoft YaHei")
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", 1)
    else:
        code_style = styles["Code Block"]
    set_style_font(code_style, "Consolas", 8.5, "111827", east_asia="Consolas")
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(7)
    code_style.paragraph_format.line_spacing = 1.0


def set_style_font(style, name: str, size: float, color: str, bold: bool = False, east_asia: str | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    if east_asia:
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    r2 = p2.add_run(subtitle)
    r2.font.name = "Calibri"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor.from_string(MUTED)


def add_metadata(doc: Document) -> None:
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["文档性质", "未来完整体 API 蓝图，不限于当前本地实现。"],
            ["当前基础", "本地已有 Temporal Python DAG workflow、Sandbox 微服务、Web 看板、simulator mock Agent、A2A 数据结构适配。"],
            ["目标场景", "大规模并发、多 Agent、多租户、长时间等待、人机回路、真实 Agent Service 接入。"],
            ["生成日期", str(date.today())],
        ],
        widths=[1.45, 5.05],
    )


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(INK)


def add_callout(doc: Document, label: str, text: str, fill: str = CALLOUT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, [6.5])
    set_table_borders(table, BORDER)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(f"{label}：")
    r1.font.bold = True
    r1.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    r1.font.name = "Calibri"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2 = p.add_run(text)
    r2.font.color.rgb = RGBColor.from_string(INK)
    r2.font.name = "Calibri"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    shade_paragraph(p, "F3F4F6")
    set_paragraph_border(p, "D1D5DB")
    for index, line in enumerate(code.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string("111827")


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table, widths)
    set_table_borders(table, BORDER)

    for cell, text in zip(table.rows[0].cells, headers):
        shade_cell(cell, HEADER_FILL)
        set_cell_margins(cell, top=90, bottom=90, start=120, end=120)
        set_cell_text(cell, text, bold=True, color=DARK_BLUE, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_values in rows:
        row = table.add_row()
        for cell, text in zip(row.cells, row_values):
            set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
            set_cell_text(cell, text, size=8.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def set_cell_text(
    cell,
    text: str,
    bold: bool = False,
    color: str = INK,
    size: float = 8.5,
    align: WD_ALIGN_PARAGRAPH | None = None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_table_width(table, widths_in: Sequence[float]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_dxa = sum(inches_to_dxa(width) for width in widths_in)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(inches_to_dxa(width)))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            set_cell_width(cell, inches_to_dxa(width))


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def inches_to_dxa(value: float) -> int:
    return int(round(value * 1440))


def set_table_borders(table, color: str) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ["top", "left", "bottom", "right"]:
        node = p_bdr.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            p_bdr.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "2")
        node.set(qn("w:color"), color)


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("持久化多 Agent 编排内核 API 蓝图")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


if __name__ == "__main__":
    main()
