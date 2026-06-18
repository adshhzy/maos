from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path.home() / "Desktop" / "持久化多Agent编排内核_API总结.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5F6B7A"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "B8C2CF"


def main() -> None:
    doc = Document()
    configure_document(doc)

    add_title(
        doc,
        "持久化多 Agent 编排内核微服务 API 总结",
        "基于 Temporal + A2A 的任务分工、编排、调度与长等待唤醒接口说明",
    )
    add_metadata(doc)
    add_callout(
        doc,
        "核心边界",
        "Sandbox 微服务负责创建和监控 DAG 任务，并把每个任务图映射为一个 Temporal Workflow；"
        "每个节点是一次 Agent 调用，节点间数据通过 A2A Message / Task / Artifact 传递；"
        "Agent 完成后回调 Sandbox，Sandbox 再 signal 对应 workflow。",
    )

    add_heading(doc, "1. API 边界总览", 1)
    add_table(
        doc,
        ["边界", "调用方向", "职责"],
        [
            [
                "Sandbox 自身 API",
                "用户 / Web UI / CLI / 上层系统 -> Sandbox",
                "创建任务、查询任务、加载示例、健康检查、接收 Agent 回调。",
            ],
            [
                "Temporal Server API",
                "Sandbox -> Temporal",
                "持久化 workflow、启动任务、列出 workflow、query 状态、signal 唤醒。",
            ],
            [
                "Agent / Simulator API",
                "Sandbox -> Agent，Agent -> Sandbox",
                "派发节点执行，Agent 完成后 callback 回 Sandbox。",
            ],
            [
                "A2A 协议契约",
                "Agent 节点之间的数据传递",
                "用 Message 表达输入，用 Task 表达执行，用 Artifact 表达上游结果。",
            ],
        ],
        widths=[1.65, 2.05, 2.8],
    )

    add_heading(doc, "2. Sandbox 自身提供的 HTTP API", 1)
    add_paragraph(doc, "当前默认 Base URL：")
    add_code(doc, "http://127.0.0.1:8765")
    add_table(
        doc,
        ["方法", "路径", "调用方", "用途"],
        [
            ["GET", "/", "浏览器", "打开 Web 可视化界面。"],
            ["POST", "/api/tasks", "Web / CLI / 上层系统", "批量创建 DAG 任务，返回 task_id 列表。"],
            ["GET", "/api/tasks", "Web / CLI / 上层系统", "查询所有任务状态，数据来自 Temporal visibility / query / result。"],
            ["GET", "/api/tasks/{task_id}", "Web / CLI / 上层系统", "查询单个任务详情，包括 DAG 状态、节点状态、结果和错误。"],
            ["POST", "/api/agent-callbacks", "Agent / Simulator", "Agent 完成或失败后回调，Sandbox 转换成 Temporal signal。"],
            ["GET", "/api/examples", "Web UI", "列出内置 JSON DAG 示例。"],
            ["GET", "/api/example?name=xxx.json", "Web UI", "读取某个内置示例 DAG JSON。"],
            ["GET", "/api/health", "运维 / 上层系统", "健康检查，返回 Sandbox、Temporal、Simulator 配置状态。"],
        ],
        widths=[0.65, 1.75, 1.65, 2.45],
    )

    add_heading(doc, "3. 核心请求与响应", 1)
    add_heading(doc, "3.1 创建任务：POST /api/tasks", 2)
    add_paragraph(doc, "请求体支持对象形式：")
    add_code(
        doc,
        '{\n'
        '  "graphs": [\n'
        '    {\n'
        '      "id": "order-processing",\n'
        '      "name": "Order processing DAG",\n'
        '      "input": {},\n'
        '      "nodes": []\n'
        '    }\n'
        '  ]\n'
        '}',
    )
    add_paragraph(doc, "响应：")
    add_code(doc, '{\n  "ok": true,\n  "task_ids": ["task-order-processing-..."]\n}')
    add_callout(
        doc,
        "实现约定",
        "返回的 task_id 就是 Temporal workflow_id。一个用户任务图对应一个独立 Temporal Workflow。",
    )

    add_heading(doc, "3.2 查询任务列表：GET /api/tasks", 2)
    add_code(
        doc,
        '{\n'
        '  "runtime_status": "ready",\n'
        '  "error": null,\n'
        '  "temporal": {\n'
        '    "mode": "embedded-persistent-dev",\n'
        '    "target": "127.0.0.1:7233",\n'
        '    "db_file": "D:\\\\dev\\\\MAOS\\\\temporal-data\\\\temporal.db",\n'
        '    "ui_port": 8233\n'
        '  },\n'
        '  "tasks": []\n'
        '}',
    )

    add_heading(doc, "3.3 查询单个任务：GET /api/tasks/{task_id}", 2)
    add_table(
        doc,
        ["字段", "说明"],
        [
            ["task_id / workflow_id", "任务 ID，同时也是 Temporal workflow_id。"],
            ["graph_id / graph_name", "DAG 图的业务标识和显示名称。"],
            ["status", "starting、running、completed、failed、cancelled、terminated、timed_out 等。"],
            ["state.workflow_status", "workflow 内部状态。"],
            ["state.levels / edges / nodes", "图结构、依赖边和所有节点状态，用于 Web 图形化展示。"],
            ["result", "workflow 完成后的最终结果；运行中为 null。"],
            ["error", "失败时的错误信息。"],
        ],
        widths=[1.8, 4.7],
    )

    add_heading(doc, "3.4 Agent 完成回调：POST /api/agent-callbacks", 2)
    add_paragraph(
        doc,
        "这个接口由 Agent 或当前 mock simulator 调用。普通用户和 Web UI 不需要直接调用。",
    )
    add_code(
        doc,
        '{\n'
        '  "workflow_id": "task-xxx",\n'
        '  "node_id": "count_items",\n'
        '  "a2a_task_id": "a2a-task-count_items-...",\n'
        '  "context_id": "task-xxx",\n'
        '  "status": "completed",\n'
        '  "job": {},\n'
        '  "result": {},\n'
        '  "error": null\n'
        '}',
    )
    add_paragraph(doc, "Sandbox 收到回调后的动作：")
    add_numbered_list(
        doc,
        [
            "将 callback 转换为 A2A Task / Artifact 结构。",
            "调用 Temporal signal：agent_node_completed(event)。",
            "唤醒正在 wait_condition 中等待该节点结果的 workflow。",
            "解锁下游依赖节点，继续 DAG 调度。",
        ],
    )

    add_heading(doc, "4. 外部依赖 API", 1)
    add_heading(doc, "4.1 Temporal Server", 2)
    add_table(
        doc,
        ["项目", "当前默认值 / 调用"],
        [
            ["Temporal Server", "127.0.0.1:7233"],
            ["Temporal UI", "http://127.0.0.1:8233"],
            ["持久 DB 文件", r"D:\dev\MAOS\temporal-data\temporal.db"],
            ["Task Queue", "json-dag-batch-visualization-task-queue"],
            ["启动 workflow", "client.start_workflow(JsonDagWorkflow.run, graph, id=task_id, task_queue=...)"],
            ["列出任务", "client.list_workflows('WorkflowType = \"JsonDagWorkflow\"')"],
            ["查询状态", "handle.query(JsonDagWorkflow.graph_state)"],
            ["获取结果", "handle.result()"],
            ["唤醒 workflow", "handle.signal(JsonDagWorkflow.agent_node_completed, event)"],
        ],
        widths=[1.8, 4.7],
    )

    add_heading(doc, "4.2 Agent / Simulator Service", 2)
    add_paragraph(doc, "当前 simulator 只是 mock Agent Service，Base URL：")
    add_code(doc, "http://127.0.0.1:8767")
    add_table(
        doc,
        ["方法", "路径", "调用方", "用途"],
        [
            ["POST", "/api/simulator/jobs", "Sandbox / A2A runtime", "创建一次模拟 Agent 执行，并提供 callback 地址。"],
            ["GET", "/api/simulator/jobs/{job_id}", "调试 / 兼容路径", "查询模拟 job 状态。当前主路径依赖 callback，不依赖轮询。"],
            ["GET", "/api/simulator/health", "运维 / Sandbox", "检查 simulator 是否在线。"],
        ],
        widths=[0.65, 1.8, 1.55, 2.5],
    )
    add_paragraph(doc, "创建模拟 Agent job 的请求体：")
    add_code(
        doc,
        '{\n'
        '  "node": {},\n'
        '  "dependency_results": {},\n'
        '  "graph_input": {},\n'
        '  "callback": {\n'
        '    "url": "http://127.0.0.1:8765/api/agent-callbacks",\n'
        '    "payload": {\n'
        '      "workflow_id": "task-xxx",\n'
        '      "node_id": "node-a",\n'
        '      "a2a_task_id": "a2a-task-node-a-...",\n'
        '      "context_id": "task-xxx"\n'
        '    }\n'
        '  }\n'
        '}',
    )

    add_heading(doc, "5. 内部协议契约", 1)
    add_heading(doc, "5.1 DAG JSON", 2)
    add_table(
        doc,
        ["字段", "说明"],
        [
            ["id", "DAG 图唯一标识。"],
            ["name", "人类可读名称。"],
            ["input", "图级输入，会传给所有节点的执行上下文。"],
            ["nodes", "节点数组，每个节点代表一次 Agent 调用。"],
        ],
        widths=[1.45, 5.05],
    )
    add_table(
        doc,
        ["节点字段", "说明"],
        [
            ["id", "节点唯一标识。"],
            ["label", "展示名称。"],
            ["operation", "当前 simulator 支持 emit、merge、template、count、percentage_from_count、status、join。"],
            ["deps", "上游节点 ID 列表。全部完成后当前节点才可运行。"],
            ["params", "节点参数，可通过 $input / $deps 引用上下文。"],
            ["timeout_seconds", "节点等待 Agent 回调的超时，默认 86400 秒。"],
            ["simulate", "mock 执行时间范围：min_seconds / max_seconds。"],
        ],
        widths=[1.45, 5.05],
    )

    add_heading(doc, "5.2 A2A 数据传递", 2)
    add_table(
        doc,
        ["A2A 元素", "在当前系统中的作用"],
        [
            ["AgentCard", "描述节点对应的 Agent 能力、接口和 skill。"],
            ["Message", "下游 Agent 的输入，包含当前 node、graph_input 和 dependency_artifacts。"],
            ["referenceTaskIds", "指向上游 Agent task，表达依赖来源。"],
            ["Task", "Agent 执行句柄，先返回 TASK_STATE_WORKING，完成后变成 TASK_STATE_COMPLETED。"],
            ["Artifact", "承载上游节点结果，下游节点从 artifact 中恢复 dependency_results。"],
        ],
        widths=[1.55, 4.95],
    )

    add_heading(doc, "6. 执行时序", 1)
    add_numbered_list(
        doc,
        [
            "用户或上层系统调用 POST /api/tasks，提交一个或多个 DAG JSON。",
            "Sandbox 校验 DAG，并为每个图启动一个 JsonDagWorkflow。",
            "Workflow 找到依赖满足的 ready nodes，并行调度这些节点。",
            "每个节点通过短 activity 调用 A2A runtime，构造 A2A Message 并创建 Agent job。",
            "Activity 很快返回 A2A task id，workflow 将节点标记为 suspended。",
            "Workflow 使用 workflow.wait_condition(...) 持久等待 Agent callback，不占用 worker 线程。",
            "Agent 或 simulator 完成后调用 POST /api/agent-callbacks。",
            "Sandbox 将 callback 转为 A2A Task / Artifact，并 signal 对应 workflow。",
            "Workflow 被唤醒，记录节点结果，解锁下游节点，直到 DAG 全部完成。",
        ],
    )

    add_heading(doc, "7. 当前部署参数", 1)
    add_table(
        doc,
        ["参数", "默认值"],
        [
            ["Sandbox Web/API", "http://127.0.0.1:8765"],
            ["Simulator API", "http://127.0.0.1:8767"],
            ["Temporal Server", "127.0.0.1:7233"],
            ["Temporal UI", "http://127.0.0.1:8233"],
            ["Temporal DB", r"D:\dev\MAOS\temporal-data\temporal.db"],
            ["自动刷新", "Web UI 每 30 秒刷新一次，也支持 Refresh Now。"],
            ["默认节点执行时间", "5 到 120 秒，可通过 simulate 覆盖。"],
        ],
        widths=[1.8, 4.7],
    )

    add_heading(doc, "8. 后续生产化关注点", 1)
    add_bullet_list(
        doc,
        [
            "把 mock simulator 替换为真实 Agent Gateway / Agent Service。",
            "把 A2A artifact 中的大 payload 改为外部对象存储引用，避免 Temporal history 膨胀。",
            "使用真实 Temporal Cluster + PostgreSQL/MySQL 替换本地 dev server。",
            "增加租户、权限、任务优先级、取消、重试策略、人工节点等业务 API。",
            "对高频进度事件做节流，只把关键状态写入 workflow history。",
        ],
    )

    add_footer(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    set_east_asian_font(normal, "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        set_east_asian_font(style, "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(30, 41, 59)
    set_east_asian_font(code, "Microsoft YaHei")
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.1


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(subtitle)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_metadata(doc: Document) -> None:
    add_table(
        doc,
        ["项目", "值"],
        [
            ["系统名称", "持久化多 Agent 编排内核微服务"],
            ["当前实现路径", r"D:\dev\MAOS\temporal_execution_core"],
            ["核心依赖", "Temporal Server、Google A2A 协议、Agent / Simulator Service"],
            ["文档用途", "API 边界、调用方式、数据契约和外部依赖说明"],
        ],
        widths=[1.45, 5.05],
    )


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [6.5])
    set_table_borders(table, BORDER)
    cell = table.cell(0, 0)
    set_cell_fill(cell, CALLOUT_FILL)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p.add_run(text)
    doc.add_paragraph()


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [6.5])
    set_table_borders(table, "D0D7E2")
    cell = table.cell(0, 0)
    set_cell_fill(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, start=140, end=140)
    p = cell.paragraphs[0]
    p.style = doc.styles["Code Block"]
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_widths(table, widths)
    set_table_borders(table, BORDER)

    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        set_cell_fill(cell, HEADER_FILL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(value))

    doc.add_paragraph()


def add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_numbered_list(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("持久化多 Agent 编排内核 API 总结")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def set_table_widths(table, widths_in_inches: list[float]) -> None:
    dxa_widths = [int(width * 1440) for width in widths_in_inches]
    table_width = sum(dxa_widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(table_width))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in dxa_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, width in enumerate(dxa_widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            tc_w.type = "dxa"
            tc_w.w = width


def set_table_borders(table, color: str) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_east_asian_font(style, font_name: str) -> None:
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


if __name__ == "__main__":
    main()
